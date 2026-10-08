import re
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from rest_framework.exceptions import ValidationError

from .sections import split_sections


SECTION_LABELS = (
    "Links", "Profile", "Professional summary", "Summary", "Objective",
    "Core strength", "Core strengths", "Core competencies",
    "Education history", "Education", "Skills & specialization",
    "Skills and specialization", "Skills", "Specialization", "Technical skills",
    "Project involvement", "Projects", "Professional experience", "Work experience",
    "Employment history", "Experience", "Certifications", "Languages", "Awards",
    "Volunteer experience", "Publications", "References",
)


def structure_resume(text):
    """Recognize common extracted labels without rewriting the applicant's facts."""
    if re.search(r"^#{1,6}\s+", text, re.M):
        return text  # An already authored Markdown document owns its structure.
    pattern = re.compile(r"^(" + "|".join(re.escape(label) for label in sorted(SECTION_LABELS, key=len, reverse=True)) + r")(?::|\s|$)", re.I)
    lines = text.splitlines()
    output = []
    index = 0
    while index < len(lines):
        line = lines[index]
        candidate = line.strip()
        consumed = 1
        # PDF label columns can wrap this single heading over two lines.
        if re.fullmatch(r"skills\s*(?:&|and)", candidate, re.I) and index + 1 < len(lines) and re.fullmatch(r"specialization", lines[index + 1].strip(), re.I):
            candidate = "Skills & specialization"
            consumed = 2
        match = pattern.match(candidate)
        if match:
            label = match[1]
            remainder = candidate[match.end():].lstrip()
            # Inline labels must be uppercase: do not mistake prose such as
            # "Experience building products" for a new section.
            if not remainder or label.isupper():
                canonical = next(item for item in SECTION_LABELS if item.casefold() == label.casefold())
                output.extend(["", f"## {canonical}", ""])
                if remainder:
                    output.append(remainder)
                index += consumed
                continue
        output.append(line)
        index += 1
    return "\n".join(output).strip()


def preserve_resume_sections(source, proposed):
    """Keep omitted sections in a review proposal, never mutate a saved draft."""
    original = split_sections(structure_resume(source))
    parts = split_sections(proposed)
    restored = []
    for index, section in enumerate(original):
        if section["title"] in ("Whole document", "Header & introduction") or not section["body"].strip():
            continue
        if any(item["key"] == section["key"] for item in parts):
            continue
        following = {item["key"] for item in original[index + 1:]}
        destination = next((position for position, item in enumerate(parts) if item["key"] in following), len(parts))
        parts.insert(destination, section)
        restored.append(section)
    if not restored:
        return proposed, []
    body = "\n\n".join(item["body"].rstrip() for item in parts)
    if len(body) > 30000:
        raise ValidationError("The proposal omitted source sections and could not retain them within 30,000 characters. Your saved résumé is unchanged. Shorten the source or draft before retrying.")
    return body, restored


def extract_resume(upload):
    if not upload or upload.size > 5 * 1024 * 1024:
        raise ValidationError({"file": "Choose a PDF, Word (.docx), or text résumé up to 5 MB."})
    extension = Path(upload.name).suffix.lower()
    raw = upload.read()
    try:
        if extension == ".txt":
            text = raw.decode("utf-8-sig")
        elif extension == ".docx":
            with ZipFile(BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20_000_000:
                    raise ValueError()
            document = Document(BytesIO(raw))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
        elif extension == ".pdf" and raw.startswith(b"%PDF-"):
            from pypdf import PdfReader
            document = PdfReader(BytesIO(raw))
            if document.is_encrypted or len(document.pages) > 30:
                raise ValueError()
            text = "\n".join(page.extract_text() or "" for page in document.pages)
        else:
            raise ValueError()
    except ImportError:
        raise ValidationError({"file": "PDF extraction is not installed yet. Install backend requirements, or upload Word/text."}) from None
    except Exception:
        # Parsers raise different errors for encrypted, damaged, and non-document uploads.
        raise ValidationError({"file": "This file could not be read. Use an unencrypted PDF, .docx, or UTF-8 text file."}) from None
    text = text.replace("\x00", "").strip()
    if not text:
        raise ValidationError({"file": "No résumé text was found. For a scanned PDF, paste the text or upload a Word file."})
    if len(text) > 30000:
        raise ValidationError({"file": "The résumé is too long. Use at most 30,000 characters."})
    text = structure_resume(text)
    if len(text) > 30000:
        raise ValidationError({"file": "The structured résumé exceeds 30,000 characters. Shorten it before uploading."})
    return text, Path(upload.name).name[:200]
