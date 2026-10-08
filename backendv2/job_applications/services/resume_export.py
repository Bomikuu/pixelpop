"""Original-style and single-column résumé exports from the same reviewed text."""
from io import BytesIO
import re
from urllib.parse import urlsplit

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor
from markdown_it import MarkdownIt

from .resume_layout import dated_entries
from .sections import split_sections


def write_inline(paragraph, children):
    bold, italic, link = False, False, None
    for token in children or []:
        if token.type in ("strong_open", "strong_close"):
            bold = token.type == "strong_open"
        elif token.type in ("em_open", "em_close"):
            italic = token.type == "em_open"
        elif token.type == "link_open":
            href = token.attrGet("href") or ""
            link = href if urlsplit(href).scheme in ("https", "http", "mailto") else None
        elif token.type == "link_close":
            link = None
        elif token.type in ("softbreak", "hardbreak"):
            paragraph.add_run("\n")
        elif token.type in ("text", "code_inline"):
            run = paragraph.add_run(token.content)
            run.bold, run.italic = bold, italic
            if link:
                hyperlink = OxmlElement("w:hyperlink")
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
                run.underline = True
                hyperlink.append(run._r)
                paragraph._p.append(hyperlink)


def render_content(document, body, heading_bold=True):
    tokens = MarkdownIt("default", {"html": False}).enable("table").parse(body)
    lists, heading, row = [], None, None
    for token in tokens:
        if token.type in ("bullet_list_open", "ordered_list_open"):
            lists.append("List Bullet" if token.type == "bullet_list_open" else "List Number")
        elif token.type in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif token.type == "heading_open":
            heading = min(int(token.tag[1:]), 6)
        elif token.type == "heading_close":
            heading = None
        elif token.type == "tr_open":
            # Flatten tables from older/manual drafts into readable text rows.
            row = document.add_paragraph()
        elif token.type == "tr_close":
            row = None
        elif token.type in ("inline", "fence", "code_block"):
            paragraph = row if row is not None else document.add_paragraph()
            if row is not None and paragraph.text:
                paragraph.add_run("; ")
            if lists:
                paragraph.style = lists[-1]
            paragraph.paragraph_format.space_after = Pt(5)
            if heading:
                paragraph.style = f"Heading {heading}"
                paragraph.paragraph_format.keep_with_next = True
            if token.type == "inline":
                write_inline(paragraph, token.children)
                if heading and (heading_bold or heading <= 2):
                    for run in paragraph.runs:
                        run.bold = True
            else:
                paragraph.add_run(token.content.rstrip("\n"))
        elif token.type == "hr":
            paragraph = document.add_paragraph()
            border = OxmlElement("w:pBdr")
            line = OxmlElement("w:bottom")
            line.set(qn("w:val"), "single")
            line.set(qn("w:sz"), "4")
            border.append(line)
            paragraph._p.get_or_add_pPr().append(border)


def fill_cell(cell, body, space_before=0):
    initial = cell.paragraphs[0]
    render_content(cell, body, heading_bold=False)
    if len(cell.paragraphs) > 1 and not initial.text:
        cell._tc.remove(initial._p)
    cell.paragraphs[0].paragraph_format.space_before = Pt(space_before)


def resume_table(container, widths, rule=False):
    table = container.add_table(rows=1, cols=len(widths))
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = table.cell(0, index).width = Mm(width)
    borders, margins = OxmlElement("w:tblBorders"), OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single" if rule and edge == "top" else "nil")
        if rule and edge == "top":
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "212121")
        borders.append(border)
        if edge in ("top", "left", "bottom", "right"):
            margin = OxmlElement(f"w:{edge}")
            margin.set(qn("w:w"), "0")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
    table._tbl.tblPr.append(borders)
    table._tbl.tblPr.append(margins)
    return table


def fill_skills(cell, body):
    tokens = MarkdownIt("default", {"html": False}).parse(body)
    simple = {"paragraph_open", "paragraph_close", "inline", "bullet_list_open", "bullet_list_close", "list_item_open", "list_item_close"}
    if any(token.type not in simple or token.level > 3 for token in tokens):
        return False
    lines = [line for token in tokens if token.type == "inline" for line in token.content.splitlines() if line.strip()]
    if len(lines) < 2:
        return False
    # The reference lists simple skills vertically, balanced across two columns.
    table = resume_table(cell, [66.75, 66.75])
    middle = (len(lines) + 1) // 2
    fill_cell(table.cell(0, 0), "\n\n".join(lines[:middle]))
    fill_cell(table.cell(0, 1), "\n\n".join(lines[middle:]))
    for index, paragraph in enumerate(cell.paragraphs):
        paragraph.paragraph_format.space_before = Pt(5 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.add_run().font.size = Pt(3)
    return True


def original_sections(document, body):
    for section in split_sections(body):
        if not section["body"].strip():
            continue
        plain = section["title"] in ("Whole document", "Header & introduction")
        text = section["body"] if plain else re.sub(r"^#{1,6}\s+[^\r\n]+(?:\r?\n)?", "", section["body"], count=1)
        label = {"Whole document": "Résumé", "Header & introduction": "Introduction"}.get(section["title"], section["title"])
        entries = dated_entries(text)
        # A thin section rule and a left label/date rail mirror the original PDF.
        table = resume_table(document, [44.5, 133.5], rule=True)
        paragraph = table.cell(0, 0).paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(label.upper())
        run.font.size, run.bold = Pt(8.75), False
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "21")
        run._r.get_or_add_rPr().append(spacing)
        if entries:
            for entry in entries:
                left, right = table.add_row().cells
                left.width, right.width = Mm(44.5), Mm(133.5)
                left.paragraphs[0].add_run(entry["date"])
                content = f"### {entry['title']}\n\n{entry['body']}" if entry["title"] else entry["body"]
                fill_cell(right, content)
                right.paragraphs[-1].paragraph_format.space_after = Pt(12)
        else:
            skills = re.fullmatch(r"skills(?:\s*&\s*speciali[sz]ation)?", section["title"].strip(), flags=re.I)
            if not skills or not fill_skills(table.cell(0, 1), text):
                fill_cell(table.cell(0, 1), text, space_before=8)
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = 1
        spacer.add_run().font.size = Pt(3)


def render_resume_docx(profile, body, layout="original"):
    original = layout == "original"
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Mm(210), Mm(297)
    page.top_margin, page.bottom_margin = Mm(10 if original else 12), Mm(15)
    page.left_margin = page.right_margin = Mm(16)
    for name in ("Normal", "List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name, style.font.size = "Times New Roman" if original else "Arial", Pt(9.5 if original else 11)
        if original:
            style.font.color.rgb = RGBColor(33, 33, 33)
        style.paragraph_format.line_spacing = 1.15 if original else 1.35
        style.paragraph_format.space_after = Pt(5)
    for level in range(1, 7):
        style = document.styles[f"Heading {level}"]
        style.font.name, style.font.size = "Times New Roman" if original else "Arial", Pt(11 if original else 12 if level <= 2 else 11)
        style.font.bold = not original or level <= 2
        style.font.color.rgb = RGBColor(33, 33, 33) if original else RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
    heading = re.match(r"^#\s+([^\r\n]+)(?:\r?\n|$)", body)
    title = heading.group(1) if heading else profile.full_name
    if heading:
        body = body[heading.end():]
    if title:
        header = document.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER if original else WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.keep_with_next = True
        if original:
            header.paragraph_format.space_after = Pt(8)
        write_inline(header, MarkdownIt().parseInline(title)[0].children)
        for run in header.runs:
            run.bold, run.font.size = True, Pt(13 if original else 16)
    contact = (", " if original else " | ").join(value for value in (profile.phone, profile.email, profile.location, profile.portfolio_url) if value)
    if contact:
        paragraph = document.add_paragraph(contact)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if original else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(16 if original else 15)
    if original:
        original_sections(document, body)
    else:
        render_content(document, body)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
