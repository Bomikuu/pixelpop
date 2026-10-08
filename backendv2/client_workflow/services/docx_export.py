from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt


def _write_inline(paragraph, children):
    bold = False
    italic = False
    for token in children or []:
        if token.type == "strong_open":
            bold = True
        elif token.type == "strong_close":
            bold = False
        elif token.type == "em_open":
            italic = True
        elif token.type == "em_close":
            italic = False
        elif token.type in ("text", "code_inline", "html_inline"):
            if token.type == "html_inline":
                continue
            run = paragraph.add_run(token.content)
            run.bold = bold
            run.italic = italic
        elif token.type in ("softbreak", "hardbreak"):
            paragraph.add_run("\n")


def render_docx(title: str, body: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    tokens = MarkdownIt("default").enable("table").parse(body)
    first_heading = tokens[1] if len(tokens) > 1 and tokens[0].type == "heading_open" and tokens[1].type == "inline" else None
    first_text = "".join(token.content for token in (first_heading.children or []) if token.type == "text") if first_heading else ""
    if first_text.strip().casefold() != title.strip().casefold():
        heading = document.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    list_stack = []
    table = None
    row = None
    cell = None
    heading_level = None
    cell_index = 0
    for index, token in enumerate(tokens):
        kind = token.type
        if kind == "bullet_list_open":
            list_stack.append("List Bullet")
        elif kind == "ordered_list_open":
            list_stack.append("List Number")
        elif kind in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
        elif kind == "heading_open":
            heading_level = min(int(token.tag[1]), 3)
        elif kind == "heading_close":
            heading_level = None
        elif kind == "table_open":
            first_row = tokens[index + 1:]
            column_count = 0
            for next_token in first_row:
                if next_token.type == "tr_close":
                    break
                if next_token.type in ("th_open", "td_open"):
                    column_count += 1
            table = document.add_table(rows=0, cols=max(column_count, 1))
            table.style = "Table Grid"
        elif kind == "table_close":
            table = None
        elif kind == "tr_open" and table is not None:
            row = table.add_row()
            cell_index = 0
        elif kind in ("th_open", "td_open") and table is not None:
            cell = row.cells[min(cell_index, len(row.cells) - 1)]
            cell_index += 1
        elif kind in ("th_close", "td_close"):
            cell = None
        elif kind == "tr_close":
            row = None
        elif kind == "hr":
            document.add_paragraph("────────────────────────")
        elif kind == "fence" or kind == "code_block":
            paragraph = document.add_paragraph(style="Normal")
            paragraph.add_run(token.content.rstrip("\n"))
        elif kind == "inline":
            if cell is not None:
                paragraph = cell.paragraphs[0]
            elif heading_level:
                paragraph = document.add_heading(level=heading_level)
                if heading_level == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif list_stack:
                paragraph = document.add_paragraph(style=list_stack[-1])
            else:
                paragraph = document.add_paragraph()
            _write_inline(paragraph, token.children)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
